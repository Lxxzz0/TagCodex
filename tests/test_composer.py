# -*- coding: utf-8 -*-
"""组合器测试：各种槽方式、预设加手写、LoRA、跨段去重、不拼接的段、导出 txt / xlsx / 组合记录、冲突、导回。临时库和临时目录。"""
import json
import os
import shutil
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import composer  # noqa: E402
import db        # noqa: E402
import gallery   # noqa: E402
import importer  # noqa: E402
import presets   # noqa: E402
import tagops    # noqa: E402


class ComposerTest(unittest.TestCase):
    def setUp(self):
        self.saved_rules = dict(importer.RULES)
        importer.RULES.update(module=None, error=None)
        self.tmp = tempfile.mkdtemp(prefix="tagcodex_composer_")
        self.conn = db.connect(os.path.join(self.tmp, "t.db"))
        db.init_db(self.conn)
        self.cfg = {"gallery_root": os.path.join(self.tmp, "图库"), "thumbs_dir": os.path.join(self.tmp, "thumbs")}
        self.a = self.load("套1", "a", ["1girl", "long hair", "bikini", "beach", "smile"])
        self.b = self.load("套2", "b", ["1girl", "school uniform", "classroom", "watermark"])
        for words, key in ((["1girl"], "character"), (["long hair"], "appearance"),
                           (["bikini", "school uniform"], "outfit"), (["beach", "classroom"], "scene")):
            tagops.assign(self.conn, [self.tid(w) for w in words], key)
        tagops.block(self.conn, [self.tid("watermark")])
        # smile 留在未归类，未归类设成不参与拼接
        tagops.patch_buckets(self.conn, {"joins": {"key": "unknown", "joins": False}})

    def tearDown(self):
        self.conn.close()
        importer.RULES.clear()
        importer.RULES.update(self.saved_rules)
        shutil.rmtree(self.tmp, ignore_errors=True)

    # ---------- 工具 ----------
    def load(self, set_name, stem, words):
        with self.conn:
            importer._import_records(self.conn, [(stem, [(None, None, words)])], set_name, "x.txt", False,
                                     importer._new_report(False))
        return self.conn.execute("SELECT id FROM images WHERE stem = ? AND set_name = ?",
                                 (stem, set_name)).fetchone()[0]

    def read(self, *parts):
        with open(os.path.join(self.tmp, *parts), encoding="utf-8") as f:
            return f.read()

    def tid(self, word):
        return self.conn.execute("SELECT id FROM tags WHERE word = ?", (word,)).fetchone()[0]

    def run_slots(self, slots, ids=None):
        rows, skipped, changed = composer.compute(self.conn, {"ids": ids or [self.a]}, slots)
        return rows, skipped, changed

    def seg(self, row, key):
        return next(s for s in row["segments"] if s["key"] == key)

    def words(self, slots, key, ids=None):
        return self.seg(self.run_slots(slots, ids)[0][0], key)["words"]

    def preset(self, name, bucket, content, lora=None):
        return presets.create(self.conn, {"name": name, "bucket": bucket, "content": content, "lora": lora})["id"]

    # ---------- 槽 ----------
    def test_keep_matches_gallery_prompt(self):
        rows, skipped, changed = self.run_slots({}, [self.a, self.b])
        segs = gallery.segments(self.conn, [self.a, self.b])
        self.assertEqual([r["prompt"] for r in rows],
                         [gallery.prompt_of(segs[self.a]), gallery.prompt_of(segs[self.b])])
        self.assertEqual((skipped, changed), (0, []))
        self.assertNotIn("watermark", rows[1]["prompt"], "黑名单词不带进来")
        self.assertEqual(self.seg(rows[1], "outfit")["words"], ["school uniform"])

    def test_replace_with_preset_and_lora(self):
        p = self.preset("角色C", "character", "ruan mei (honkai: star rail), 1girl", "ruanmei:0.8")
        rows, _, changed = self.run_slots({"character": {"mode": "replace", "presets": [p], "text": "solo"}})
        s = self.seg(rows[0], "character")
        self.assertEqual(s["words"], ["<lora:ruanmei:0.8>", "ruan mei (honkai: star rail)", "1girl", "solo"],
                         "LoRA 在最前，预设在手写前")
        self.assertEqual(s["added"], ["<lora:ruanmei:0.8>", "ruan mei (honkai: star rail)", "solo"])
        self.assertEqual(changed, ["character"])
        self.assertTrue(rows[0]["prompt"].startswith("<lora:ruanmei:0.8>, ruan mei"))

    def test_prepend_append_clear(self):
        p = self.preset("x", "scene", "sunset", "sky:1")
        self.assertEqual(self.words({"scene": {"mode": "prepend", "text": "night"}}, "scene"), ["night", "beach"])
        self.assertEqual(self.words({"scene": {"mode": "append", "presets": [p], "text": "night"}}, "scene"),
                         ["<lora:sky:1>", "beach", "sunset", "night"], "后加时 LoRA 也在这段最前")
        self.assertEqual(self.words({"scene": {"mode": "clear"}}, "scene"), [])
        self.assertEqual(self.words({"scene": {"mode": "keep", "text": "ignored"}}, "scene"), ["beach"])

    def test_remove_uses_normalize_and_alias(self):
        self.load("套3", "c", ["two piece"])
        tagops.merge_alias(self.conn, self.tid("bikini"), [self.tid("two piece")])
        self.assertEqual(self.words({"outfit": {"mode": "remove", "text": "Two_Piece"}}, "outfit"), [])
        s = self.seg(self.run_slots({"scene": {"mode": "remove", "text": "BEACH"}})[0][0], "scene")
        self.assertEqual((s["words"], s["removed"]), ([], ["beach"]))

    def test_exclude_preset_words_only_in_slot(self):
        p = self.preset("角色C", "character", "ruan mei (honkai: star rail), 1girl, Solo", "ruanmei:0.8")
        slot = {"mode": "replace", "presets": [p], "text": "solo, smile", "exclude": ["1GIRL", "solo", " "]}
        self.assertEqual(self.words({"character": slot}, "character"),
                         ["<lora:ruanmei:0.8>", "ruan mei (honkai: star rail)", "solo", "smile"],
                         "去掉的词按规范化比较，只跳过预设里的，手写的不受影响")
        self.assertEqual(presets.get_preset(self.conn, p)["content"], "ruan mei (honkai: star rail), 1girl, Solo",
                         "不改预设")
        self.assertEqual(composer.resolve_slots(self.conn, {"character": slot})["character"]["exclude"],
                         ["1GIRL", "solo"])
        r = self.export(slots={"character": slot})
        rec = json.loads(self.read("out", composer.RECORD_FILE))
        self.assertEqual(rec["slots"]["character"]["exclude"], ["1GIRL", "solo"], "组合记录里也记上")

    def test_dedupe_across_segments(self):
        w = self.words({"scene": {"mode": "append", "text": "1girl, Sunset, sunset"}}, "scene")
        self.assertEqual(w, ["beach", "Sunset"], "角色段已有 1girl；段内重复也去掉")

    def test_nojoin_segment_kept_out_of_prompt(self):
        row = self.run_slots({})[0][0]
        self.assertEqual(self.seg(row, "unknown")["words"], ["smile"])
        self.assertNotIn("smile", row["prompt"])
        self.assertIn("[未归类]\nsmile\n", composer.txt_of(row["segments"]))

    def test_slot_errors(self):
        p = self.preset("x", "scene", "sunset")
        for slots in ({"character": {"mode": "replace", "presets": [p]}}, {"nope": {"mode": "keep"}},
                      {"scene": {"mode": "bad"}}):
            with self.assertRaises(tagops.OpError):
                self.run_slots(slots)
        with self.assertRaises(tagops.OpError):
            composer.compute(self.conn, {}, {})

    def test_source_filter_and_skipped(self):
        with self.conn:
            empty = self.conn.execute("INSERT INTO images (path, stem, set_name, added) VALUES "
                                      "('import:套1/e', 'e', '套1', '2026-10-11')").lastrowid
        rows, skipped, _ = composer.compute(self.conn, {"filter": {"set": "套1", "sort": "name"}}, {})
        self.assertEqual(([r["image_id"] for r in rows], skipped), ([self.a], 1))
        rows, skipped, _ = self.run_slots({}, [self.b, empty, self.a])
        self.assertEqual([r["image_id"] for r in rows], [self.b, self.a], "按给的顺序")
        page = composer.preview(self.conn, {"source": {"ids": [self.a, self.b]}, "slots": {}, "size": 1})
        self.assertEqual((page["total"], len(page["items"]), page["skipped"]), (2, 1, 0))

    # ---------- 导出 ----------
    def export(self, **kw):
        body = {"source": {"ids": kw.pop("ids", [self.a, self.b])}, "slots": kw.pop("slots", {}),
                "name": "批次", "dir": os.path.join(self.tmp, "out"), "txt": True, "xlsx": True}
        body.update(kw)
        return composer.export(self.conn, self.cfg, body)

    def test_source_stats(self):
        st = composer.source_stats(self.conn, {"source": {"ids": [self.a, self.b]}})
        segs = {x["key"]: x for x in st["segments"]}
        self.assertEqual(st["total"], 2)
        self.assertEqual([(t["word"], t["count"], t["ratio"]) for t in segs["character"]["tags"]], [("1girl", 2, 1.0)])
        self.assertEqual(sorted(t["word"] for t in segs["outfit"]["tags"]), ["bikini", "school uniform"])
        same = gallery.stats(self.conn, {"set": "套1"})
        self.assertEqual(composer.source_stats(self.conn, {"source": {"filter": {"set": "套1"}}}), same,
                         "按筛选和图库词频一致")
        with self.assertRaises(tagops.OpError):
            composer.source_stats(self.conn, {"source": None})

    def test_export_files_and_conflict(self):
        dup = self.load("套9", "a", ["1girl", "beach"])
        r = self.export(ids=[self.a, self.b, dup])
        out = r["dir"]
        self.assertEqual(sorted(os.listdir(out)), sorted(["a.txt", "b.txt", "a_2.txt", "批次.xlsx", composer.RECORD_FILE]))
        rec = json.loads(self.read("out", composer.RECORD_FILE))
        self.assertEqual([x["file"] for x in rec["images"]], ["a.txt", "b.txt", "a_2.txt"])
        # 有同名文件：不写任何文件
        before = self.read("out", "a.txt")
        with self.assertRaises(composer.Conflict) as cm:
            self.export(slots={"scene": {"mode": "clear"}})
        self.assertIn("a.txt", cm.exception.files)
        self.assertEqual(self.read("out", "a.txt"), before)
        self.export(slots={"scene": {"mode": "clear"}}, overwrite=True)
        self.assertIn("[环境]\n\n", self.read("out", "a.txt"))

    def test_export_xlsx_layout(self):
        from openpyxl import load_workbook
        self.export(txt=False)
        out = os.path.join(self.tmp, "out")
        self.assertFalse(os.path.exists(os.path.join(out, "a.txt")))
        ws = load_workbook(os.path.join(out, "批次.xlsx")).worksheets[0]
        head = [c.value for c in ws[1]]
        names = [r["name"] for r in self.conn.execute("SELECT name FROM buckets ORDER BY ord")]
        self.assertEqual(head, ["序号", "图", "文件名"] + names + ["完整提示词"])
        self.assertEqual((ws.title, ws.max_row, ws["C2"].value), (composer.SHEET_TITLE, 3, "a"))

    def test_export_dir_rules(self):
        with self.assertRaises(tagops.OpError):
            self.export(dir=os.path.join(self.cfg["gallery_root"], "区", "out"))
        with self.assertRaises(tagops.OpError):
            self.export(dir=" ")
        with self.assertRaises(tagops.OpError):
            self.export(txt=False, xlsx=False)

    def test_roundtrip_import(self):
        p = self.preset("角色C", "character", "ruan mei (honkai: star rail)", "ruanmei:0.8")
        self.export(slots={"character": {"mode": "prepend", "presets": [p]}}, dir=os.path.join(self.tmp, "导回"))
        rows = composer.compute(self.conn, {"ids": [self.a, self.b]},
                                {"character": {"mode": "prepend", "presets": [p]}})[0]
        report = importer.import_txt_dir(self.conn, os.path.join(self.tmp, "导回"))
        self.assertEqual(report["lora_skipped"], 2, "LoRA 标记不建成词")
        self.assertIsNone(self.conn.execute("SELECT 1 FROM tags WHERE word LIKE '<lora%'").fetchone())
        for stem, row in zip(("a", "b"), rows):
            iid = self.conn.execute("SELECT id FROM images WHERE set_name = '导回' AND stem = ?", (stem,)).fetchone()[0]
            got = {s["key"]: [t["word"] for t in s["tags"]] for s in gallery.segments(self.conn, [iid])[iid]}
            want = {s["key"]: [w for w in s["words"] if not w.startswith("<lora")] for s in row["segments"]}
            # 新词 ruan mei 按机器归类落段（这里是未归类），所以只比词的全集和没加新词的段
            got_words = [w for k in got for w in got[k]]
            self.assertEqual(sorted(got_words), sorted(w for k in want for w in want[k]))
            for key in ("appearance", "outfit", "scene"):
                self.assertEqual(got[key], want[key], key)


class LoraTest(unittest.TestCase):
    def test_normalize_lora(self):
        self.assertEqual(presets.normalize_lora("ruanmei:0.8\n<lora:a:b:0.5>, x"),
                         "<lora:ruanmei:0.8>\n<lora:a:b:0.5>\n<lora:x:1>")
        self.assertIsNone(presets.normalize_lora(" \n"))
        with self.assertRaises(tagops.OpError):
            presets.normalize_lora("<bad>")


if __name__ == "__main__":
    unittest.main()
