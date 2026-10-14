# -*- coding: utf-8 -*-
"""
组合器：拿一批来源图按段分好的提示词，按槽批量替换，预览，导出 txt 目录 / xlsx；
没有来源时按槽算单条提示词（compose）。

来源图的分段按词典实时算（gallery.segments），黑名单词不带进来；拼提示词用 gallery.prompt_of。
只读数据库；导出只写用户指定的输出目录（不能在图库根里面），不进撤销栈。
"""
import io
import json
import os
import re

import db
import gallery
import importer
import scanner
import tagops

MAX_SOURCE = 1000
# auto：有词替换、没词保留（默认）；keep / replace 是旧槽的写法，照认
MODES = ("auto", "keep", "replace", "prepend", "append", "remove", "clear")
MODE_NAMES = {"keep": "保留", "replace": "替换", "prepend": "前加", "append": "后加", "remove": "删词", "clear": "清空"}
RECORD_FILE = "_组合记录.json"
SHEET_TITLE = "组合结果"
# xlsx 样式：表头底色、字色、缩略图尺寸
HEAD_FILL = "2F5597"
THUMB_W = 90
COL_W = {"序号": 6, "图": 14, "文件名": 24, "完整提示词": 60}
SEG_COL_W = 30
ROW_H_NO_THUMB = 90


class Conflict(tagops.OpError):
    """输出目录里已有同名文件、又没让覆盖；files 是冲突的文件名。"""

    def __init__(self, files):
        super().__init__("输出目录里已经有 %d 个同名文件" % len(files))
        self.files = files


# ---------------------------------------------------------------------
#  来源与槽
# ---------------------------------------------------------------------
def source_ids(conn, source):
    """{"ids": [...]} 按给的顺序；{"filter": {...}} 按图库当时的筛选和排序展开。"""
    source = source or {}
    if source.get("ids") is not None:
        ids = gallery._image_ids(source["ids"])
        found = set()
        for chunk in range(0, len(ids), 500):           # SQLite 参数个数有上限
            part = ids[chunk:chunk + 500]
            found |= {r[0] for r in conn.execute("SELECT id FROM images WHERE id IN (%s)" % ",".join("?" * len(part)),
                                                 part)}
        ids = [i for i in ids if i in found]
    elif isinstance(source.get("filter"), dict):
        sql, params, order = gallery._filters(conn, source["filter"])
        ids = [r[0] for r in conn.execute("SELECT id FROM (%s) i ORDER BY %s" % (sql, order), params)]
    else:
        raise tagops.OpError("来源要么是勾选的图（ids），要么是图库筛选（filter）")
    if not ids:
        raise tagops.OpError("来源里没有图")
    if len(ids) > MAX_SOURCE:
        raise tagops.OpError("一批最多 %d 张，来源有 %d 张，先缩小范围" % (MAX_SOURCE, len(ids)))
    return ids


def source_stats(conn, body):
    """{source} → 来源图每段的词频，格式同 gallery.stats（预设组编辑器的候选词）。id 列表用 json_each 传，不受参数个数限制。"""
    ids = source_ids(conn, (body or {}).get("source"))
    return gallery.stats_of(conn, "SELECT value AS id FROM json_each(?)", [json.dumps(ids)])


def _split(text):
    return [w for w in (" ".join(x.split()) for x in re.split(r"[,\n]", text or "")) if w]


def resolve_slots(conn, slots):
    """
    {key: {mode, text}} → {key: {mode, words, text}}；没写的段按保留。auto 按有没有词换成 replace / keep。
    旧槽里的 presets / exclude 忽略（预设改成以文本填进段）。组合器不处理 LoRA。
    """
    keys = [r["key"] for r in conn.execute("SELECT key FROM buckets ORDER BY ord")]
    out = {}
    for key, s in (slots or {}).items():
        if key not in keys:
            raise tagops.OpError("段不存在：%s" % key)
        s = s or {}
        mode = s.get("mode") or "auto"
        if mode not in MODES:
            raise tagops.OpError("不认识的方式：%s" % mode)
        text = s.get("text") or ""
        words = _split(text) if mode not in ("keep", "clear") else []
        if mode == "auto":
            mode = "replace" if words else "keep"
        out[key] = {"mode": mode, "words": words, "text": text if mode != "keep" else ""}
    return out


def _key(word):
    """去重和删词的比较键：LoRA 标记原样比，其余按规范化后的写法。"""
    return word if importer.LORA_RE.match(word) else importer.normalize(word)


def _apply(orig, slot, remove_keys):
    mode = slot["mode"]
    if mode == "keep":
        return list(orig)
    if mode == "clear":
        return []
    if mode == "replace":
        return list(slot["words"])
    if mode == "prepend":
        return slot["words"] + orig
    if mode == "append":
        return orig + slot["words"]
    return [w for w in orig if _key(w) not in remove_keys]       # remove


# ---------------------------------------------------------------------
#  计算
# ---------------------------------------------------------------------
def compute(conn, source, slots):
    """
    → (rows, skipped, changed_segments)。rows 按来源顺序，每行：
      {image_id, stem, path, set_name, thumb, segments: [{key, name, joins, words, added, removed, text}], prompt}
    没有提示词的来源图跳过（skipped 计数）；空白套图的条（drafts）没词也不跳过。
    """
    ids = source_ids(conn, source)
    slots = resolve_slots(conn, slots)
    remove_keys = _remove_keys(conn, slots)
    segs = gallery.segments(conn, ids)
    info = {}
    for chunk in range(0, len(ids), 500):
        part = ids[chunk:chunk + 500]
        for r in conn.execute("SELECT id, path, stem, set_name, sha1, missing, origin FROM images WHERE id IN (%s)"
                              % ",".join("?" * len(part)), part):
            info[r["id"]] = dict(r)
    rows, skipped, changed = [], 0, set()
    for i in ids:
        src = [{"key": s["key"], "name": s["name"], "joins": s["joins"],
                "orig": [t["word"] for t in s["tags"] if not t["blocked"]]} for s in segs[i]]
        if not any(s["orig"] for s in src) and info[i]["origin"] != "draft":
            skipped += 1
            continue
        out = _compose(src, slots, remove_keys)
        changed |= {s["key"] for s, o in zip(src, out) if o["words"] != s["orig"]}
        im = info[i]
        placeholder = im["path"].startswith("import:")
        rows.append({"image_id": i, "stem": im["stem"], "path": im["path"], "set_name": im["set_name"],
                     "missing": bool(im["missing"]), "placeholder": placeholder,
                     "thumb": None if placeholder else (im["sha1"] or "")[:8] or "0",
                     "segments": out, "prompt": prompt_of(out)})
    order = [r["key"] for r in conn.execute("SELECT key FROM buckets ORDER BY ord")]
    return rows, skipped, [k for k in order if k in changed]


def _remove_keys(conn, slots):
    """删词按主词比：手写的别名也认"""
    return {k: {_key(importer._resolve_alias(conn, importer.normalize(w))) for w in s["words"]}
            for k, s in slots.items() if s["mode"] == "remove"}


def _compose(src, slots, remove_keys):
    """一张图按段套槽：src [{key, name, joins, orig}] → [{key, name, joins, words, added, removed, text}]。"""
    seen, out = set(), []
    for s in src:
        slot = slots.get(s["key"], {"mode": "keep"})
        words = []
        for w in _apply(s["orig"], slot, remove_keys.get(s["key"], set())):
            k = _key(w)
            if k and k not in seen:                   # 整条提示词里同一个词只留第一次出现
                seen.add(k)
                words.append(w)
        before = {_key(w) for w in s["orig"]}
        after = {_key(w) for w in words}
        out.append({"key": s["key"], "name": s["name"], "joins": s["joins"], "words": words,
                    "added": [w for w in words if _key(w) not in before],
                    "removed": [w for w in s["orig"] if _key(w) not in after],
                    "text": gallery.PROMPT_SEP.join(words)})
    return out


def compose(conn, body):
    """
    单条编辑（没有来源）：把槽套到一张空图上 → {segments, prompt, all_text}。
    prompt 跳过不参与拼接的段；all_text 含全部段，给页面文本框同步用（再拆一遍不丢词）。
    """
    slots = resolve_slots(conn, body.get("slots"))
    src = [{"key": b["key"], "name": b["name"], "joins": bool(b["joins"]), "orig": []} for b in gallery._buckets(conn)]
    out = _compose(src, slots, _remove_keys(conn, slots))
    segs = [{k: s[k] for k in ("key", "name", "joins", "words", "text")} for s in out]
    return {"segments": segs, "prompt": prompt_of(out),
            "all_text": gallery.PROMPT_SEP.join(w for s in out for w in s["words"])}


def prompt_of(segs):
    """拼法交给 gallery.prompt_of（跳过不拼接的段），这里只换成它要的结构。"""
    return gallery.prompt_of([{"joins": s["joins"], "tags": [{"word": w, "blocked": False} for w in s["words"]]}
                              for s in segs])


def preview(conn, body):
    rows, skipped, changed = compute(conn, body.get("source"), body.get("slots"))
    page = db.paginate_list(rows, body.get("page", 1), body.get("size", 60))
    page.update(skipped=skipped, changed_segments=changed)
    return page


# ---------------------------------------------------------------------
#  导出
# ---------------------------------------------------------------------
def safe_name(name):
    """批次名 / 文件名里不能有的字符换成 _。"""
    return re.sub(r'[\\/:*?"<>|]+', "_", str(name or "")).strip(" .") or "untitled"


def file_names(rows):
    """每行一个 txt 文件名主干：同一批里重名的从第二个起加 _2、_3。"""
    used, out = set(), []
    for r in rows:
        base = safe_name(r["stem"] or r["image_id"])
        name, n = base, 1
        while name.lower() in used:
            n += 1
            name = "%s_%d" % (base, n)
        used.add(name.lower())
        out.append(name)
    return out


def txt_of(segs):
    """分段 txt：每段一节（一行 [段名]、下一行这段的词），空段也写节名，顺序和名称跟词典；导入时按同样的格式读回。"""
    lines = []
    for s in segs:
        lines += ["[%s]" % s["name"], ", ".join(s["words"]), ""]
    return "\n".join(lines).rstrip() + "\n"


def _check_dir(cfg, out_dir):
    out_dir = (out_dir or "").strip().strip('"')
    if not out_dir:
        raise tagops.OpError("要填输出目录")
    out_dir = os.path.abspath(out_dir)
    root = cfg.get("gallery_root")
    if root:
        root = os.path.abspath(root)
        try:
            inside = os.path.commonpath([os.path.normcase(root), os.path.normcase(out_dir)]) == os.path.normcase(root)
        except ValueError:                                # 不同盘符
            inside = False
        if inside:
            raise tagops.OpError("输出目录不能在图库根目录里面：%s" % root)
    return out_dir


def export(conn, cfg, body):
    """
    {source, slots, name, dir, txt, xlsx, overwrite} → {dir, files, count, skipped}。
    dir 就是最终目录（前端默认拼好 composer_output_dir/<批次名>）。有同名文件又没让覆盖时抛 Conflict，什么都不写。
    """
    want_txt, want_xlsx = bool(body.get("txt", True)), bool(body.get("xlsx", True))
    if not want_txt and not want_xlsx:
        raise tagops.OpError("txt 和 xlsx 至少导出一种")
    name = safe_name(body.get("name"))
    out_dir = _check_dir(cfg, body.get("dir"))
    rows, skipped, changed = compute(conn, body.get("source"), body.get("slots"))
    if not rows:
        raise tagops.OpError("来源图都没有提示词，没有可导出的")
    stems = file_names(rows)
    plan = ([s + ".txt" for s in stems] if want_txt else []) + ([name + ".xlsx"] if want_xlsx else []) + [RECORD_FILE]
    conflicts = [f for f in plan if os.path.exists(os.path.join(out_dir, f))]
    if conflicts and not body.get("overwrite"):
        raise Conflict(conflicts)
    try:
        os.makedirs(out_dir, exist_ok=True)
        if want_txt:
            for stem, r in zip(stems, rows):
                with open(os.path.join(out_dir, stem + ".txt"), "w", encoding="utf-8") as f:
                    f.write(txt_of(r["segments"]))
        if want_xlsx:
            write_xlsx(cfg, os.path.join(out_dir, name + ".xlsx"), rows, stems)
        _write_record(conn, out_dir, name, body, rows, stems)
    except OSError as e:
        raise tagops.OpError("写文件失败：%s" % e)
    return {"dir": out_dir, "files": plan, "count": len(rows), "skipped": skipped, "changed_segments": changed}


def _thumb_png(cfg, row):
    """缩略图 webp → 宽 THUMB_W 的 png（BytesIO）；没有图返回 None。"""
    if row["placeholder"]:
        return None
    path = scanner.thumb_path(cfg, row["image_id"])
    if not os.path.exists(path):
        if row["missing"] or not os.path.isfile(row["path"]):
            return None
        path = scanner.ensure_thumb(cfg, row["image_id"], row["path"])
        if path is None:
            return None
    from PIL import Image
    try:
        with Image.open(path) as im:
            im = im.convert("RGB")
            h = max(1, int(im.height * THUMB_W / im.width))
            buf = io.BytesIO()
            im.resize((THUMB_W, h), Image.LANCZOS).save(buf, format="PNG")
    except OSError:
        return None
    buf.seek(0)
    return buf


def write_xlsx(cfg, path, rows, stems):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    seg_names = [s["name"] for s in rows[0]["segments"]]
    heads = ["序号", "图", "文件名"] + seg_names + ["完整提示词"]
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    fill, font = PatternFill("solid", fgColor=HEAD_FILL), Font(color="FFFFFF", bold=True, size=11)
    widths = [COL_W["序号"], COL_W["图"], COL_W["文件名"]] + [SEG_COL_W] * len(seg_names) + [COL_W["完整提示词"]]
    for ci, (title, width) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=1, column=ci, value=title)
        c.fill, c.font = fill, font
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "D2"
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center")
    for idx, (stem, r) in enumerate(zip(stems, rows)):
        row = idx + 2
        ws.cell(row=row, column=1, value=idx + 1).alignment = center
        ws.cell(row=row, column=3, value=stem).alignment = wrap
        for j, s in enumerate(r["segments"]):
            ws.cell(row=row, column=4 + j, value=", ".join(s["words"])).alignment = wrap
        ws.cell(row=row, column=len(heads), value=r["prompt"]).alignment = wrap
        png = _thumb_png(cfg, r)
        if png is None:
            ws.row_dimensions[row].height = ROW_H_NO_THUMB
            continue
        xi = XLImage(png)
        ws.row_dimensions[row].height = max(60, xi.height * 0.75)
        ws.add_image(xi, "B%d" % row)
    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)


def _write_record(conn, out_dir, name, body, rows, stems):
    slots = {}
    for key, s in resolve_slots(conn, body.get("slots")).items():
        if s["mode"] == "keep":
            continue
        slots[key] = {"mode": s["mode"], "mode_name": MODE_NAMES[s["mode"]], "text": s["text"]}
    record = {"name": name, "time": db.now(), "dir": out_dir, "slots": slots,
              "images": [{"id": r["image_id"], "path": r["path"], "file": stem + ".txt"}
                         for stem, r in zip(stems, rows)]}
    with open(os.path.join(out_dir, RECORD_FILE), "w", encoding="utf-8") as f:
        json.dump(record, f, ensure_ascii=False, indent=1)
